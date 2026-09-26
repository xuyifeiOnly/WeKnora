"""Regression tests for non-conforming styles.xml fills (#3637).

Real-world exporters (cloud-doc systems, LabView, Acumatica ERP) write
``<fill>`` elements openpyxl cannot deserialize. The fixtures below replace
the first fill of an openpyxl-written workbook at the zip level, which
reproduces the production failure without an original file.
"""

import io
import os
import re
import subprocess
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch

import openpyxl
import pandas as pd

from docreader.parser import xlsx_repair
from docreader.parser.excel_parser import _prepare_xlsx_bytes
from docreader.parser.markitdown_parser import StdMarkitdownParser
from docreader.parser.xlsx_merge import fill_merged_cells_xlsx
from docreader.parser.xlsx_repair import sanitize_xlsx_styles

CRASHING_FILLS = {
    "empty-self-closing": b"<fill/>",
    "capitalized": b"<Fill/>",
    "text-content": b"<fill>oops</fill>",
}

# openpyxl happens to tolerate an unknown fill child, but it is still
# non-conforming; the sanitizer normalizes it like the crashing forms.
TOLERATED_FILL = b"<fill><Foo/></fill>"


def _workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "hello"
    ws["B2"] = 42
    ws["C3"] = "world"
    ws.merge_cells("C3:D3")
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _with_first_fill_replaced(workbook: bytes, fragment: bytes) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(workbook)) as zin, zipfile.ZipFile(
        out, "w", zipfile.ZIP_DEFLATED
    ) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/styles.xml":
                data = re.sub(
                    b"<fill>.*?</fill>", fragment, data, count=1, flags=re.S
                )
            zout.writestr(item, data)
    return out.getvalue()


def _with_rich_styles_part(workbook: bytes) -> bytes:
    """Add the xl/richData/richStyles.xml part Excel 365 writes for in-cell
    images / data types; its name also ends in ``styles.xml``."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(workbook)) as zin, zipfile.ZipFile(
        out, "w", zipfile.ZIP_DEFLATED
    ) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr(
            "xl/richData/richStyles.xml",
            b'<richStyleSheet xmlns="http://schemas.microsoft.com/office/'
            b'spreadsheetml/2017/richdata2"/>',
        )
    return out.getvalue()


class TestSanitizeXlsxStyles(unittest.TestCase):
    def test_clean_workbook_is_returned_unchanged(self):
        self.assertIsNone(sanitize_xlsx_styles(_workbook_bytes()))

    def test_non_zip_input_is_ignored(self):
        self.assertIsNone(sanitize_xlsx_styles(b"not a zip"))

    def test_repair_is_logged_with_count(self):
        fixture = _with_first_fill_replaced(
            _workbook_bytes(), CRASHING_FILLS["empty-self-closing"]
        )
        with self.assertLogs("docreader.parser.xlsx_repair", level="INFO") as logs:
            sanitize_xlsx_styles(fixture)
        self.assertTrue(
            any("Sanitized 1 non-conforming fill" in line for line in logs.output),
            logs.output,
        )

    def test_internal_failure_returns_none_instead_of_raising(self):
        fixture = _with_first_fill_replaced(
            _workbook_bytes(), CRASHING_FILLS["empty-self-closing"]
        )
        with patch(
            "docreader.parser.xlsx_repair._rewrite_zip", side_effect=OSError("boom")
        ):
            with self.assertLogs("docreader.parser.xlsx_repair", level="WARNING"):
                result = sanitize_xlsx_styles(fixture)
        self.assertIsNone(result)

    def test_each_malformed_form_is_repaired_in_place(self):
        for name, fragment in {**CRASHING_FILLS, "tolerated-unknown-child": TOLERATED_FILL}.items():
            with self.subTest(form=name):
                fixture = _with_first_fill_replaced(_workbook_bytes(), fragment)

                sanitized = sanitize_xlsx_styles(fixture)
                self.assertIsNotNone(sanitized)

                wb = openpyxl.load_workbook(io.BytesIO(sanitized), data_only=True)
                ws = wb.active
                self.assertEqual(ws["A1"].value, "hello")
                self.assertEqual(ws["B2"].value, 42)

                with zipfile.ZipFile(io.BytesIO(sanitized)) as z:
                    styles = z.read("xl/styles.xml")
                count = re.search(rb'<fills count="(\d+)"', styles)
                num_fills = len(re.findall(rb"<fill[ />]", styles))
                self.assertIsNotNone(count)
                self.assertEqual(int(count.group(1)), num_fills)

    def test_workbook_stylesheet_is_picked_over_rich_styles_part(self):
        fixture = _with_rich_styles_part(
            _with_first_fill_replaced(
                _workbook_bytes(), CRASHING_FILLS["empty-self-closing"]
            )
        )
        sanitized = sanitize_xlsx_styles(fixture)
        self.assertIsNotNone(sanitized)
        openpyxl.load_workbook(io.BytesIO(sanitized), data_only=True)

    def test_stylesheet_choice_does_not_depend_on_hash_seed(self):
        # Zip member names are str, so any set-ordered lookup changes with
        # PYTHONHASHSEED; run the sanitizer under several seeds.
        fixture = _with_rich_styles_part(
            _with_first_fill_replaced(
                _workbook_bytes(), CRASHING_FILLS["empty-self-closing"]
            )
        )
        repo_root = os.path.dirname(
            os.path.dirname(os.path.dirname(os.path.abspath(xlsx_repair.__file__)))
        )
        script = (
            "import sys\n"
            "from docreader.parser.xlsx_repair import sanitize_xlsx_styles\n"
            "data = open(sys.argv[1], 'rb').read()\n"
            "sys.exit(0 if sanitize_xlsx_styles(data) is not None else 1)\n"
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            path = os.path.join(tmpdir, "rich.xlsx")
            with open(path, "wb") as handle:
                handle.write(fixture)
            for seed in range(8):
                with self.subTest(seed=seed):
                    env = {
                        **os.environ,
                        "PYTHONHASHSEED": str(seed),
                        "PYTHONPATH": os.pathsep.join(
                            filter(None, [repo_root, os.environ.get("PYTHONPATH")])
                        ),
                    }
                    result = subprocess.run(
                        [sys.executable, "-c", script, path], env=env
                    )
                    self.assertEqual(result.returncode, 0)

    def test_crashing_forms_raise_without_the_fix(self):
        for name, fragment in CRASHING_FILLS.items():
            with self.subTest(form=name):
                fixture = _with_first_fill_replaced(_workbook_bytes(), fragment)
                with self.assertRaises(TypeError):
                    openpyxl.load_workbook(io.BytesIO(fixture), data_only=True)


class TestBuiltinChain(unittest.TestCase):
    def test_prepare_xlsx_bytes_recovers_and_keeps_merge_fill(self):
        for name, fragment in CRASHING_FILLS.items():
            with self.subTest(form=name):
                fixture = _with_first_fill_replaced(_workbook_bytes(), fragment)
                with self.assertRaises(TypeError):
                    fill_merged_cells_xlsx(fixture)  # the unfixed failure

                prepared = _prepare_xlsx_bytes(fixture)
                excel = pd.ExcelFile(io.BytesIO(prepared), engine="openpyxl")
                self.assertTrue(excel.sheet_names)

                wb = openpyxl.load_workbook(io.BytesIO(prepared), data_only=True)
                ws = wb.active
                # the merge-cell prefill still applied on the sanitized bytes
                self.assertEqual(ws["C3"].value, "world")
                self.assertEqual(ws["D3"].value, "world")


class TestMarkitdownChain(unittest.TestCase):
    def test_markitdown_engine_parses_sanitized_workbook(self):
        for name, fragment in {**CRASHING_FILLS, "tolerated-unknown-child": TOLERATED_FILL}.items():
            with self.subTest(form=name):
                fixture = _with_first_fill_replaced(_workbook_bytes(), fragment)
                parser = StdMarkitdownParser(file_type="xlsx")
                doc = parser.parse_into_text(fixture)
                self.assertIn("hello", doc.content)
                self.assertIn("world", doc.content)


if __name__ == "__main__":
    unittest.main()
